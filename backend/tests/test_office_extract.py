"""Office / structured-document extractors + image normalization.

Generates small real files with the same libs the extractors use, so this
verifies the wiring end-to-end (import path, dispatch, content) without shipping
binary fixtures. Legacy .doc (antiword) and .heic (a real camera file) are
covered by the in-container validation, not here.
"""

from __future__ import annotations

import io

import pytest

from app.intake import ocr


def test_docx_extraction(tmp_path):
    from docx import Document

    d = Document()
    d.add_paragraph("Service connection for tinnitus is granted.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Diagnostic Code"
    t.rows[0].cells[1].text = "6260"
    p = tmp_path / "note.docx"
    d.save(p)
    out = ocr._extract_office(str(p), ".docx")
    assert "tinnitus" in out.lower()
    assert "6260" in out


def test_rtf_extraction(tmp_path):
    p = tmp_path / "letter.rtf"
    p.write_text(r"{\rtf1\ansi Dr. Fermo notes chronic knee pain since service.\par}")
    out = ocr._extract_office(str(p), ".rtf")
    assert "knee pain" in out.lower()


def test_xlsx_extraction(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Vaccine", "Date"])
    ws.append(["Hepatitis B", "2010-02-06"])
    p = tmp_path / "immunizations.xlsx"
    wb.save(p)
    out = ocr._extract_office(str(p), ".xlsx")
    assert "hepatitis b" in out.lower()
    assert "2010-02-06" in out


def test_csv_extraction(tmp_path):
    p = tmp_path / "labs.csv"
    p.write_text("test,value\nLyme,negative\n")
    out = ocr._extract_office(str(p), ".csv")
    assert "lyme" in out.lower()


def test_office_dispatch_never_raises_on_garbage(tmp_path):
    p = tmp_path / "broken.docx"
    p.write_bytes(b"not really a docx")
    assert ocr._extract_office(str(p), ".docx") == ""


def test_image_normalization_bmp_to_png():
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (12, 12), "white").save(buf, format="BMP")
    png = ocr._normalize_image_to_png(buf.getvalue())
    assert png is not None
    assert png[:8] == b"\x89PNG\r\n\x1a\n"  # PNG magic


def test_image_normalization_returns_none_on_nonimage():
    assert ocr._normalize_image_to_png(b"definitely not an image") is None


def test_new_extensions_are_accepted():
    from app.routers.documents import ALLOWED_EXTS

    for ext in (".heic", ".bmp", ".webp", ".docx", ".doc", ".rtf", ".odt", ".xlsx", ".xls", ".csv"):
        assert ext in ALLOWED_EXTS, f"{ext} should be accepted"


@pytest.mark.parametrize("ext", [".docx", ".doc", ".rtf", ".odt", ".xlsx", ".xls", ".csv"])
def test_office_exts_route_through_office_branch(ext):
    assert ext in ocr._OFFICE_EXTS
